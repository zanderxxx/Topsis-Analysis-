import pandas as pd
from pathlib import Path
import json

from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill


def export_report(
    result_df,
    log_df,
    output_excel,
    screen_order=None
):
    chart_df = _chart_dataframe(
        result_df,
        screen_order
    )

    # -------------------------
    # Excel 输出
    # -------------------------
    with pd.ExcelWriter(
        output_excel,
        engine="openpyxl"
    ) as writer:

        result_df.to_excel(
            writer,
            index=False,
            sheet_name="TOPSIS关注度排名"
        )

        log_df.to_excel(
            writer,
            index=False,
            sheet_name="异常修正日志"
        )

        chart_df.to_excel(
            writer,
            index=False,
            sheet_name="关注度得分图表"
        )

        _format_workbook(
            writer.book,
            chart_df
        )

    # -------------------------
    # JSON 输出
    # -------------------------
    json_path = (
        Path(output_excel)
        .with_suffix(".json")
    )

    result_df.to_json(
        json_path,
        orient="records",
        force_ascii=False,
        indent=2
    )

    print(f"JSON 已导出: {json_path}")

    # -------------------------
    # 交互式 HTML 图表输出
    # -------------------------
    html_path = (
        Path(output_excel)
        .with_name(
            Path(output_excel).stem +
            "_interactive_chart.html"
        )
    )

    _export_interactive_chart(
        chart_df,
        html_path
    )

    print(f"交互式图表已导出: {html_path}")


def _chart_dataframe(
    result_df,
    screen_order
):
    chart_df = result_df[
        [
            "每一屏名称",
            "TOPSIS关注度得分",
            "关注度排名"
        ]
    ].copy()

    if screen_order:
        order_map = {
            screen: idx
            for idx, screen in enumerate(screen_order)
        }

        chart_df["_screen_order"] = (
            chart_df["每一屏名称"]
            .map(order_map)
            .fillna(len(order_map))
        )

        chart_df = (
            chart_df
            .sort_values(
                by=[
                    "_screen_order",
                    "每一屏名称"
                ]
            )
            .drop(columns=["_screen_order"])
        )

    return chart_df


def _format_workbook(
    workbook,
    chart_df
):
    header_fill = PatternFill(
        "solid",
        fgColor="1F4E79"
    )
    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for sheet_name in [
        "TOPSIS关注度排名",
        "异常修正日志",
        "关注度得分图表"
    ]:
        sheet = workbook[sheet_name]
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        sheet.freeze_panes = "A2"

    chart_sheet = workbook["关注度得分图表"]
    chart_sheet.column_dimensions["A"].width = 34
    chart_sheet.column_dimensions["B"].width = 18
    chart_sheet.column_dimensions["C"].width = 12

    for cell in chart_sheet["B"][1:]:
        cell.number_format = "0.0000"

    chart = LineChart()
    chart.title = "每一屏 TOPSIS 关注度得分趋势（按屏幕顺序）"
    chart.y_axis.title = "TOPSIS关注度得分"
    chart.x_axis.title = "每一屏名称"
    chart.height = 12
    chart.width = 24

    row_count = len(chart_df) + 1
    data = Reference(
        chart_sheet,
        min_col=2,
        min_row=1,
        max_row=row_count
    )
    categories = Reference(
        chart_sheet,
        min_col=1,
        min_row=2,
        max_row=row_count
    )
    chart.add_data(
        data,
        titles_from_data=True
    )
    chart.set_categories(categories)

    chart_sheet.add_chart(
        chart,
        "E2"
    )


def _export_interactive_chart(
    chart_df,
    html_path
):
    data_rows = [
        {
            "screen": row["每一屏名称"],
            "score": float(row["TOPSIS关注度得分"]),
            "rank": int(row["关注度排名"])
        }
        for _, row in chart_df.iterrows()
    ]

    points_js = ",\n      ".join(
        "{"
        f"screen: {json.dumps(row['screen'], ensure_ascii=False)}, "
        f"score: {row['score']:.4f}, "
        f"rank: {row['rank']}"
        "}"
        for row in data_rows
    )

    html = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>各屏 TOPSIS 关注度得分趋势</title>
  <style>
    body {{
      margin: 0;
      padding: 28px;
      background: #f7f9fc;
      color: #1f2937;
      font: 14px/1.45 -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    }}
    main {{
      max-width: 1280px;
      margin: 0 auto;
      background: #fff;
      border: 1px solid #dbe3ee;
      border-radius: 8px;
      padding: 24px;
      box-shadow: 0 12px 30px rgba(15, 23, 42, 0.08);
    }}
    h1 {{
      margin: 0 0 8px;
      font-size: 22px;
      letter-spacing: 0;
    }}
    p {{
      margin: 0 0 20px;
      color: #64748b;
    }}
    svg {{
      display: block;
      width: 100%;
      height: auto;
      overflow: visible;
    }}
    .axis {{
      fill: #64748b;
      font-size: 11px;
    }}
    .axis-line {{
      stroke: #bac6d4;
    }}
    .grid {{
      stroke: #d7dee8;
      stroke-dasharray: 4 4;
    }}
    .series-line {{
      fill: none;
      stroke: #2f6fb6;
      stroke-width: 2.5;
    }}
    .point {{
      fill: #fff;
      stroke: #2f6fb6;
      stroke-width: 2;
    }}
    .hit {{
      fill: transparent;
      cursor: pointer;
    }}
    #tooltip {{
      position: fixed;
      z-index: 10;
      pointer-events: none;
      display: none;
      min-width: 170px;
      padding: 10px 12px;
      border-radius: 6px;
      background: rgba(15, 23, 42, 0.92);
      color: #fff;
      box-shadow: 0 12px 28px rgba(15, 23, 42, 0.24);
      font-size: 13px;
    }}
    #tooltip strong {{
      display: block;
      margin-bottom: 4px;
      word-break: break-word;
    }}
    #tooltip span {{
      display: block;
      color: #dbeafe;
    }}
  </style>
</head>
<body>
  <main>
    <h1>各屏 TOPSIS 关注度得分趋势（按屏幕顺序）</h1>
    <p>鼠标悬停到折线节点上查看该屏幕的排名和值。</p>
    <svg id="chart" viewBox="0 0 1180 560" role="img" aria-label="各屏 TOPSIS 关注度得分趋势折线图"></svg>
  </main>
  <div id="tooltip"></div>
  <script>
    const data = [
      {points_js}
    ];
    const svg = document.getElementById("chart");
    const tooltip = document.getElementById("tooltip");
    const width = 1180;
    const height = 560;
    const margin = {{ top: 34, right: 32, bottom: 160, left: 64 }};
    const innerWidth = width - margin.left - margin.right;
    const innerHeight = height - margin.top - margin.bottom;
    const x = (index) => margin.left + (data.length === 1 ? innerWidth / 2 : index * innerWidth / (data.length - 1));
    const y = (score) => margin.top + innerHeight - score * innerHeight;
    const make = (tag, attrs = {{}}) => {{
      const node = document.createElementNS("http://www.w3.org/2000/svg", tag);
      for (const [key, value] of Object.entries(attrs)) node.setAttribute(key, value);
      svg.appendChild(node);
      return node;
    }};

    for (let tick = 0; tick <= 5; tick += 1) {{
      const value = tick / 5;
      const yy = y(value);
      make("line", {{ x1: margin.left, y1: yy, x2: width - margin.right, y2: yy, class: "grid" }});
      const text = make("text", {{ x: margin.left - 12, y: yy + 4, "text-anchor": "end", class: "axis" }});
      text.textContent = value.toFixed(4);
    }}

    make("line", {{ x1: margin.left, y1: margin.top, x2: margin.left, y2: margin.top + innerHeight, class: "axis-line" }});
    make("line", {{ x1: margin.left, y1: margin.top + innerHeight, x2: width - margin.right, y2: margin.top + innerHeight, class: "axis-line" }});

    const pathData = data.map((d, i) => `${{i === 0 ? "M" : "L"}} ${{x(i).toFixed(2)}} ${{y(d.score).toFixed(2)}}`).join(" ");
    make("path", {{ d: pathData, class: "series-line" }});

    const showTooltip = (event, d) => {{
      tooltip.style.display = "block";
      tooltip.innerHTML = `<strong>${{d.screen}}</strong><span>排名 ${{d.rank}}（${{d.score.toFixed(4)}}）</span>`;
      const offset = 14;
      tooltip.style.left = `${{Math.min(event.clientX + offset, window.innerWidth - tooltip.offsetWidth - 12)}}px`;
      tooltip.style.top = `${{Math.min(event.clientY + offset, window.innerHeight - tooltip.offsetHeight - 12)}}px`;
    }};
    const hideTooltip = () => {{
      tooltip.style.display = "none";
    }};

    data.forEach((d, i) => {{
      const xx = x(i);
      const yy = y(d.score);
      const label = make("text", {{
        x: xx,
        y: margin.top + innerHeight + 18,
        transform: `rotate(-90 ${{xx}} ${{margin.top + innerHeight + 18}})`,
        "text-anchor": "end",
        class: "axis",
      }});
      label.textContent = d.screen;
      make("circle", {{ cx: xx, cy: yy, r: 4.5, class: "point" }});
      const hit = make("circle", {{ cx: xx, cy: yy, r: 12, class: "hit" }});
      hit.addEventListener("mouseenter", (event) => showTooltip(event, d));
      hit.addEventListener("mousemove", (event) => showTooltip(event, d));
      hit.addEventListener("mouseleave", hideTooltip);
    }});
  </script>
</body>
</html>
"""

    html_path.write_text(
        html,
        encoding="utf-8"
    )
